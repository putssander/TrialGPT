import json
import sys
import os
from docx import Document  # you may need to install python-docx via pip

eps = 1e-9

def get_matching_score(matching):
    # count only the valid ones
    included = 0
    not_inc = 0
    na_inc = 0
    no_info_inc = 0

    excluded = 0
    not_exc = 0
    na_exc = 0
    no_info_exc = 0

    exclusion_parse_error = False

    # first count inclusions
    try:
        for criteria, info in matching["inclusion"].items():
            if len(info) != 3:
                continue

            if info[2] == "included":
                included += 1    
            elif info[2] == "not included":
                not_inc += 1
            elif info[2] == "not applicable":
                na_inc += 1
            elif info[2] == "not enough information":
                no_info_inc += 1
    except AttributeError:
        # If matching["inclusion"] is not a dict of items, treat as not enough information
        no_info_inc += len(matching.get("inclusion", {})) if isinstance(matching.get("inclusion"), dict) else 1

    # then count exclusions
    try:
        for criteria, info in matching["exclusion"].items():
            if len(info) != 3:
                continue

            if info[2] == "excluded":
                excluded += 1    
            elif info[2] == "not excluded":
                not_exc += 1
            elif info[2] == "not applicable":
                na_exc += 1
            elif info[2] == "not enough information":
                no_info_exc += 1
            # Example: if a "skipped" value is encountered, treat it as not enough information
            elif info[2] == "skipped":
                no_info_exc += 1
    except AttributeError:
        exclusion_parse_error = True
        no_info_exc += len(matching.get("exclusion", {})) if isinstance(matching.get("exclusion"), dict) else 1

    # get the matching score
    score = 0
    score += included / (included + not_inc + no_info_inc + eps)
    if not_inc > 0:
        score -= 1
    if excluded > 0:
        score -= 1

    return score, exclusion_parse_error


def get_agg_score(assessment):
    try:
        rel_score = float(assessment["relevance_score_R"])
        eli_score = float(assessment["eligibility_score_E"])
    except:
        rel_score = 0
        eli_score = 0

    score = (rel_score + eli_score) / 100
    return score 


if __name__ == "__main__":
    # args are the results paths
    matching_results_path = sys.argv[1]
    agg_results_path = sys.argv[2]

    # loading the results
    matching_results = json.load(open(matching_results_path))
    agg_results = json.load(open(agg_results_path))

    # define a top-level results folder and create if it doesn't exist
    results_folder = "results"
    os.makedirs(results_folder, exist_ok=True)

    # loop over the patients
    for patient_id, label2trial2results in matching_results.items():
        trial2info = {}

        for _, trial2results in label2trial2results.items():
            for trial_id, results in trial2results.items():
                matching_score, excl_error = get_matching_score(results)

                if patient_id not in agg_results or trial_id not in agg_results[patient_id]:
                    print(f"Patient {patient_id} Trial {trial_id} not in the aggregation results.")
                    agg_score = 0
                else:
                    agg_score = get_agg_score(agg_results[patient_id][trial_id])

                trial_score = matching_score + agg_score
                trial2info[trial_id] = {
                    "score": trial_score,
                    "exclusion_parse_error": excl_error
                }

        # sort trials by score descending
        sorted_trial2info = sorted(trial2info.items(), key=lambda x: -x[1]["score"])

        # Create a subfolder for the patient
        patient_folder = os.path.join(results_folder, f"patient_{patient_id}")
        os.makedirs(patient_folder, exist_ok=True)

        # JSON output per patient including error information
        json_out_path = os.path.join(patient_folder, f"patient_{patient_id}_ranking.json")
        with open(json_out_path, "w") as json_file:
            json.dump(sorted_trial2info, json_file, indent=4)

        # DOCX output per patient with an extra column for the error flag
        doc = Document()
        doc.add_heading(f"Patient ID: {patient_id}", level=1)
        doc.add_paragraph("Clinical trial ranking:")

        # Create a table with three columns: Trial, Score, Exclusion Parse Error
        table = doc.add_table(rows=1, cols=3)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Trial"
        hdr_cells[1].text = "Score"
        hdr_cells[2].text = "Exclusion Parse Error"
        for trial, info in sorted_trial2info:
            row_cells = table.add_row().cells
            row_cells[0].text = str(trial)
            row_cells[1].text = str(info["score"])
            row_cells[2].text = "Yes" if info["exclusion_parse_error"] else "No"

        docx_out_path = os.path.join(patient_folder, f"patient_{patient_id}_ranking.docx")
        doc.save(docx_out_path)

        # Additionally, dump key parts of the input files for this patient
        # For example, store the original matching results for this patient.
        input_summary_path = os.path.join(patient_folder, f"patient_{patient_id}_input_summary.json")
        with open(input_summary_path, "w") as summary_file:
            json.dump(label2trial2results, summary_file, indent=4)

        # Print results to console
        print()
        print(f"Patient ID: {patient_id}")
        print("Clinical trial ranking:")
        for trial, info in sorted_trial2info:
            print(trial, info["score"], "Exclusion Error:", "Yes" if info["exclusion_parse_error"] else "No")
        print("===")
        print()